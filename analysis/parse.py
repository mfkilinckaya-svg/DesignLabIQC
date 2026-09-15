import openpyxl, re, os, glob, datetime, csv

HDR_LOOSE = re.compile(r'Sysmex\s*XR\s*(\d)', re.I)
LVL = re.compile(r'CHECK\s*(L\d)', re.I)
LOT = re.compile(r'\[(\d+)\]')
INC = re.compile(r'Inkludert:\s*(\d+)', re.I)
EXC = re.compile(r'Ekskludert:\s*(\d+)', re.I)
SUMHDR = 'Beregnet gjennomsnitt'

def num(s):
    if s is None: return None
    if isinstance(s,(int,float)): return float(s)
    m = re.search(r'-?\d+[.,]?\d*', str(s))
    return float(m.group(0).replace(',','.')) if m else None

def parse_range(s):
    if not s: return (None,None)
    t = str(s).replace(',', '.')
    m = re.search(r'(\d+\.?\d*)\s*[-–—]\s*(\d+\.?\d*)', t)
    if not m: return (None,None)
    return (float(m.group(1)), float(m.group(2)))

rows, summaries = [], []

for path in sorted(glob.glob('data/Data/*/*.xlsx')):
    folder = os.path.basename(os.path.dirname(path))
    fname = os.path.basename(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    for sh in wb.sheetnames:
        ws = wb[sh]
        comp_title, cur = None, None
        pending_sum = False
        for r in range(1, ws.max_row+1):
            a = ws.cell(r,1).value
            if isinstance(a,str) and a.strip()=='Inkludert' and ws.cell(r,3).value==SUMHDR:
                pending_sum = True; continue
            if pending_sum:
                pending_sum = False
                if cur:
                    inc = ws.cell(r,1).value
                    if isinstance(inc, datetime.datetime):
                        inc = (inc - datetime.datetime(1899,12,30)).days
                    summaries.append(dict(
                        folder=folder, file=fname, instrument=cur['inst'],
                        level=cur['level'], lot=cur['lot'],
                        inkludert=inc, ekskludert=num(ws.cell(r,2).value),
                        mean_beregnet=num(ws.cell(r,3).value),
                        mean_angitt=num(ws.cell(r,4).value),
                        sd_beregnet=num(ws.cell(r,5).value),
                        sd_angitt=num(ws.cell(r,6).value),
                        cv_beregnet=num(ws.cell(r,7).value),
                        cv_angitt=num(ws.cell(r,8).value),
                        zstat=num(ws.cell(r,9).value)))
                continue
            if isinstance(a,str):
                if a.strip().startswith('Slice-sammendrag'):
                    cur = None; continue
                if HDR_LOOSE.search(a) and 'CHECK' in a.upper():
                    m = HDR_LOOSE.search(a); lm = LVL.search(a); lt = LOT.search(a)
                    inc = INC.search(a); exc = EXC.search(a)
                    cur = dict(inst='XR'+m.group(1),
                               level=lm.group(1).upper() if lm else None,
                               lot=lt.group(1) if lt else None,
                               inkludert=int(inc.group(1)) if inc else None,
                               ekskludert=int(exc.group(1)) if exc else None)
                    continue
                if comp_title is None and a.strip() and not HDR_LOOSE.search(a) and 'Verifisert' not in a:
                    comp_title = a.strip()
            if isinstance(a, datetime.datetime) and cur and a.year >= 2000:
                nxt = ws.cell(r+1,1).value
                rule = nxt if isinstance(nxt,str) and 'S-' in nxt else None
                lo, hi = parse_range(ws.cell(r,9).value)
                rows.append(dict(
                    folder=folder, file=fname, component=(ws.cell(r,6).value or comp_title),
                    tidspunkt=a.isoformat(sep=' '), instrument=cur['inst'],
                    level=cur['level'], lot=cur['lot'],
                    verdi=num(ws.cell(r,8).value), ref_lo=lo, ref_hi=hi,
                    regel=rule, avvist=(1 if rule and 'Mislyktes' in rule else 0) if rule else None))

for name, data in (('parsed_raw.csv', rows), ('beaker_summary.csv', summaries)):
    with open(name,'w',newline='',encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(data[0].keys())); w.writeheader(); w.writerows(data)

print('measurements:', len(rows))
print('summary blocks:', len(summaries))
